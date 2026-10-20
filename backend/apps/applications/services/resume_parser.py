import io
import re
from pathlib import Path
from django.core.exceptions import ValidationError
import pypdf
import docx

MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024  # 5 MB
ALLOWED_EXTENSIONS = {".pdf", ".docx"}

# Magic byte signatures
MAGIC_BYTES_PDF = b"%PDF"
MAGIC_BYTES_DOCX = b"PK\x03\x04"

# Standard technical skill taxonomy for reliable keyword matching
STANDARD_SKILL_TAXONOMY = [
    # Languages
    "Python", "JavaScript", "TypeScript", "Java", "C++", "C#", "C", "Go", "Golang",
    "Ruby", "PHP", "Swift", "Kotlin", "Rust", "Scala", "R", "SQL", "HTML", "CSS",
    # Frameworks & Libraries
    "Django", "Django REST Framework", "Flask", "FastAPI", "React", "React Native",
    "Next.js", "Vue", "Angular", "Node.js", "Express", "Spring Boot", "ASP.NET",
    "Laravel", "Tailwind CSS", "Bootstrap", "Redux", "GraphQL", "jQuery",
    # Databases & Storage
    "PostgreSQL", "MySQL", "SQLite", "MongoDB", "Redis", "Elasticsearch",
    "Cassandra", "DynamoDB", "Oracle", "Firebase",
    # Cloud & DevOps
    "AWS", "Amazon Web Services", "Azure", "GCP", "Google Cloud", "Docker",
    "Kubernetes", "Terraform", "CI/CD", "GitHub Actions", "GitLab CI", "Jenkins",
    "Linux", "Nginx", "Apache", "Ansible", "Prometheus", "Grafana",
    # Concepts & Methodologies
    "REST", "RESTful APIs", "Microservices", "System Design", "Agile", "Scrum",
    "TDD", "Unit Testing", "Machine Learning", "Deep Learning", "NLP", "Pandas",
    "NumPy", "Scikit-Learn", "PyTorch", "TensorFlow", "Git", "GitHub", "Jira",
]

DEGREE_KEYWORDS = [
    "Ph.D.", "PhD", "Doctorate", "Master", "M.S.", "M.Sc.", "M.Tech", "MBA",
    "Bachelor", "B.S.", "B.Sc.", "B.Tech", "B.E.", "Associate Degree", "Diploma",
]


def validate_resume_file(uploaded_file):
    """
    Validate uploaded resume file extension, size cap, and header signatures.
    Raises ValidationError if any condition is violated.
    """
    filename = uploaded_file.name
    ext = Path(filename).suffix.lower()

    if ext not in ALLOWED_EXTENSIONS:
        raise ValidationError(
            f"Unsupported file format '{ext}'. Only PDF and DOCX files are accepted."
        )

    if uploaded_file.size > MAX_FILE_SIZE_BYTES:
        max_mb = MAX_FILE_SIZE_BYTES / (1024 * 1024)
        raise ValidationError(
            f"File size exceeds the {max_mb:.0f} MB limit (uploaded size: {uploaded_file.size / (1024 * 1024):.2f} MB)."
        )

    # Check magic byte signatures
    uploaded_file.seek(0)
    header = uploaded_file.read(4)
    uploaded_file.seek(0)  # Reset cursor for downstream read

    if ext == ".pdf" and not header.startswith(MAGIC_BYTES_PDF):
        raise ValidationError("Invalid PDF file: Missing standard PDF file signature.")

    if ext == ".docx" and not header.startswith(MAGIC_BYTES_DOCX):
        raise ValidationError("Invalid DOCX file: Corrupted or missing OpenXML ZIP container signature.")


def extract_text_from_pdf(file_stream):
    """
    Extract text content from PDF stream using pypdf.
    """
    text_chunks = []
    reader = pypdf.PdfReader(file_stream)
    for idx, page in enumerate(reader.pages):
        page_text = page.extract_text()
        if page_text:
            text_chunks.append(page_text.strip())
    return "\n\n".join(text_chunks)


def extract_text_from_docx(file_stream):
    """
    Extract text content from DOCX stream using python-docx.
    """
    doc = docx.Document(file_stream)
    text_chunks = []

    # Extract paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            text_chunks.append(para.text.strip())

    # Extract table cells
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                text_chunks.append(" | ".join(row_text))

    return "\n".join(text_chunks)


def extract_entities_from_text(raw_text):
    """
    Deterministic extraction of contact info, skills, and education.
    Strictly reports only facts present in the text; never hallucinates or invents data.
    """
    if not raw_text or not raw_text.strip():
        return {
            "emails": [],
            "phones": [],
            "skills": [],
            "education": [],
            "word_count": 0,
        }

    # 1. Email extraction (RFC 5322 compatible regex subset)
    email_regex = r"\b[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+\b"
    emails = list(dict.fromkeys(re.findall(email_regex, raw_text, re.IGNORECASE)))

    # 2. Phone number extraction (common formats with optional international prefix)
    phone_regex = r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}\b"
    raw_phones = re.findall(phone_regex, raw_text)
    phones = list(dict.fromkeys([p.strip() for p in raw_phones if len(re.sub(r"\D", "", p)) >= 10]))

    # 3. Technical Skills matching against standard taxonomy
    detected_skills = []
    text_lower = raw_text.lower()
    for skill in STANDARD_SKILL_TAXONOMY:
        # Match using word boundaries to prevent substring false positives (e.g. 'C' inside 'CSS')
        pattern = r"\b" + re.escape(skill.lower()) + r"\b"
        if re.search(pattern, text_lower):
            detected_skills.append(skill)

    # 4. Education indicators
    detected_education = []
    for deg in DEGREE_KEYWORDS:
        pattern = r"\b" + re.escape(deg.lower()) + r"\b"
        if re.search(pattern, text_lower):
            detected_education.append(deg)

    return {
        "emails": emails,
        "phones": phones,
        "skills": detected_skills,
        "education": detected_education,
        "word_count": len(raw_text.split()),
    }


def parse_resume_instance(resume_instance):
    """
    Orchestrates reading the stored resume file, extracting raw text,
    running deterministic entity detection, and persisting results on the Resume model.
    Handles malformed or unreadable documents gracefully without throwing exceptions.
    """
    from apps.applications.models import ParsingStatus

    try:
        ext = Path(resume_instance.original_filename).suffix.lower()

        # Open file stream from Django storage
        with resume_instance.file.open("rb") as f:
            file_bytes = io.BytesIO(f.read())

        if ext == ".pdf":
            raw_text = extract_text_from_pdf(file_bytes)
        elif ext == ".docx":
            raw_text = extract_text_from_docx(file_bytes)
        else:
            raise ValueError(f"Unsupported extension: {ext}")

        if not raw_text.strip():
            resume_instance.parsing_status = ParsingStatus.FAILED
            resume_instance.parsing_error = "Document contained no extractable text (e.g., scanned image or empty file)."
            resume_instance.raw_text = ""
            resume_instance.parsed_data = {}
        else:
            entities = extract_entities_from_text(raw_text)
            resume_instance.raw_text = raw_text
            resume_instance.parsed_data = entities
            resume_instance.parsing_status = ParsingStatus.PARSED
            resume_instance.parsing_error = ""

    except Exception as exc:
        resume_instance.parsing_status = ParsingStatus.FAILED
        resume_instance.parsing_error = f"Extraction failed: {str(exc)}"
        resume_instance.raw_text = ""
        resume_instance.parsed_data = {}

    resume_instance.save(update_fields=["raw_text", "parsed_data", "parsing_status", "parsing_error"])
    return resume_instance
